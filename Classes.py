import os
import sys
import re
import time
from datetime import datetime
from docx import Document
import comtypes.client
from PyQt5 import QtCore, QtWidgets, QtGui
from PyQt5.QtCore import pyqtSlot
from PyQt5.QtWidgets import QApplication, QDialog, QPushButton, QWidget, QMainWindow
from PyQt5.uic import loadUi



global client_name
global doa
global attorney_name
global user_name
global user_email
global user_fax
global user_phone
global output_path
global widget
global lor_pl
global lop_pl
global rr_pl
global doc

#Path to files

path = os.getcwd()
settings_path = (path + "/settings")
lop_pl = (path + "\placeholders\LoP.docx")
rr_pl = (path + "\placeholders\RR.docx")
lor_pl = (path + "\placeholders\LoR.docx")
user_fax_pl = re.compile("User Fax")
user_email_pl = re.compile("User Email")
current_balance_pl = re.compile("Current Balance")
requested_balance_pl = re.compile("Requested Balance")

os.chdir (path+r"\gui")

def tabFocus(self):
    self.plainTextEdit.setTabChangesFocus(True)
    self.plainTextEdit_2.setTabChangesFocus(True)
    self.plainTextEdit_3.setTabChangesFocus(True)
    self.plainTextEdit_4.setTabChangesFocus(True)
    self.plainTextEdit_5.setTabChangesFocus(True)
    self.plainTextEdit_6.setTabChangesFocus(True)
    self.plainTextEdit_7.setTabChangesFocus(True)
    self.plainTextEdit_8.setTabChangesFocus(True)
    self.plainTextEdit_9.setTabChangesFocus(True)
    self.setTabOrder(self.plainTextEdit, self.plainTextEdit_2)
    self.setTabOrder(self.plainTextEdit_2, self.plainTextEdit_3)
    self.setTabOrder(self.plainTextEdit_3, self.plainTextEdit_4)
    self.setTabOrder(self.plainTextEdit_4, self.plainTextEdit_5)
    self.setTabOrder(self.plainTextEdit_5, self.plainTextEdit_6)
    self.setTabOrder(self.plainTextEdit_6, self.plainTextEdit_7)
    self.setTabOrder(self.plainTextEdit_7, self.plainTextEdit_8)
    self.setTabOrder(self.plainTextEdit_8, self.plainTextEdit_9)




class Window(QDialog):
    def pagesetup(self, currentpage):
        self.currentpage = currentpage
        super(currentpage, self).__init__()
        loadUi("mainpage"+".ui", self)

    def cancel_click(self):
        sys.exit

    def move(self, target):
        self.target.show()
        self.hide
        doc_type = target


class mainpage(Window):
    def __init__(self):
        super(mainpage, self).__init__()
        Window.pagesetup(self, mainpage)
        pushButton.clicked.Window.move(self, settings)




#if os.path.exists(settings_path + "/user_name.txt") and os.path.exists(settings_path + "/user_email.txt") and os.path.exists(settings_path + "/user_phone.txt") and os.path.exists(settings_path + "/user_email.txt"):
app = QApplication(sys.argv)
global widget
widget = mainpage()
widget.show()
sys.exit(app.exec_())

#else:
#    app = QApplication(sys.argv)
#    widget = settings()
#    widget.show()
#    sys.exit(app.exec_())